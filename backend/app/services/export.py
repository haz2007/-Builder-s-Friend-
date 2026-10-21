"""Выгрузка заказов в Excel: для учёта ИП и сверки с кассой.

Два листа: «Заказы» (один заказ — одна строка, итог формулой) и
«Позиции» (что именно продано). Время — московское, как в магазине.
"""

import io
from datetime import UTC, datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session, selectinload

from ..models import Order, OrderItem
from .order_notifications import status_name

MOSCOW = timezone(timedelta(hours=3))
HEADER = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="3A2A1F")
MONEY = '#,##0.00 "₽"'


def moscow_time(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    if value.tzinfo is None:
        value = value.replace(tzinfo=UTC)
    # Excel не хранит часовой пояс — пишем местное время без него
    return value.astimezone(MOSCOW).replace(tzinfo=None)


def _sheet(workbook: Workbook, title: str, columns: list[tuple[str, int]], first: bool = False):
    sheet = workbook.active if first else workbook.create_sheet()
    assert sheet is not None
    sheet.title = title
    sheet.append([name for name, _ in columns])
    for index, (_, width) in enumerate(columns, start=1):
        cell = sheet.cell(row=1, column=index)
        cell.font, cell.fill = HEADER, HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    return sheet


def orders_workbook(db: Session, start: datetime | None, end: datetime | None) -> bytes:
    query = db.query(Order).options(
        selectinload(Order.items).selectinload(OrderItem.product),
        selectinload(Order.items).selectinload(OrderItem.variant),
        selectinload(Order.user),
    )
    if start is not None:
        query = query.filter(Order.created_at >= start)
    if end is not None:
        query = query.filter(Order.created_at < end)
    orders = query.order_by(Order.id).all()

    workbook = Workbook()
    summary = _sheet(workbook, "Заказы", [
        ("№", 7), ("Дата", 17), ("Статус", 18), ("Покупатель", 20), ("Телефон", 15),
        ("Получение", 12), ("Адрес", 36), ("Комментарий", 30), ("Сумма", 14),
    ], first=True)
    lines = _sheet(workbook, "Позиции", [
        ("№ заказа", 9), ("Дата", 17), ("Товар", 36), ("Вариант", 14),
        ("Кол-во", 9), ("Ед.", 7), ("Цена", 12), ("Сумма", 14), ("Статус", 18),
    ])

    for order in orders:
        created = moscow_time(order.created_at)
        status = status_name(order.status, order.delivery_method)
        pickup = order.delivery_method == "PICKUP"
        summary.append([
            order.id, created, status, order.customer_name, order.customer_phone,
            "Самовывоз" if pickup else "Доставка", "" if pickup else order.delivery_address,
            order.comment or "", order.total_price,
        ])
        for item in order.items:
            lines.append([
                order.id, created,
                getattr(item.product, "name", None) or f"Товар №{item.product_id}",
                getattr(item.variant, "name", None) or item.unit,
                item.quantity, item.unit, item.price, round(item.price * item.quantity, 2), status,
            ])

    for sheet, date_column, money_columns in ((summary, "B", "I"), (lines, "B", "GH")):
        for row in range(2, sheet.max_row + 1):
            sheet[f"{date_column}{row}"].number_format = "DD.MM.YYYY HH:MM"
            for column in money_columns:
                sheet[f"{column}{row}"].number_format = MONEY

    # Итог формулой — без отменённых; Excel пересчитает, если удалить строки
    last = summary.max_row
    total_row = last + 2
    summary[f"H{total_row}"] = "Итого без отменённых:"
    summary[f"H{total_row}"].font = Font(bold=True)
    summary[f"I{total_row}"] = f'=SUMIF(C2:C{last},"<>Отменён",I2:I{last})'
    summary[f"I{total_row}"].font = Font(bold=True)
    summary[f"I{total_row}"].number_format = MONEY
    summary.auto_filter.ref = f"A1:I{last}"

    data = io.BytesIO()
    workbook.save(data)
    return data.getvalue()
