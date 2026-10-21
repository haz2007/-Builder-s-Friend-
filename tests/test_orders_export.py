"""Выгрузка заказов в Excel для учёта."""

import io
from datetime import datetime

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.auth import create_access_token, hash_password
from app.database import SessionLocal
from app.models import Category, Order, OrderItem, Product, User


def setup_orders():
    db = SessionLocal()
    admin = User(email="admin@woodcommerce.local", password_hash=hash_password("x" * 8), role="ADMIN")
    customer = User(phone="+79991112233", password_hash=hash_password("x" * 8), role="CUSTOMER")
    category = Category(name="Доска")
    db.add_all([admin, customer, category])
    db.flush()
    product = Product(name="Доска 50×150", price=620, stock_quantity=100, unit="шт", category_id=category.id)
    db.add(product)
    db.flush()
    for status, quantity in (("DELIVERED", 3), ("CANCELLED", 10)):
        order = Order(
            user_id=customer.id, customer_name="Сергей", status=status, total_price=620 * quantity,
            delivery_method="PICKUP", delivery_address="Самовывоз со склада", comment="распил по 3 м",
            created_at=datetime(2026, 10, 20, 9, 30),
        )
        db.add(order)
        db.flush()
        db.add(OrderItem(order_id=order.id, product_id=product.id, variant_id=product.variants[0].id,
                         quantity=quantity, price=620, unit="шт", base_quantity=1))
    db.commit()
    headers = {"Authorization": f"Bearer {create_access_token(user_id=admin.id, role='ADMIN')}"}
    db.close()
    return headers


def test_admin_downloads_orders_as_excel(client: TestClient):
    headers = setup_orders()

    response = client.get("/api/v1/admin/orders/export?period=all", headers=headers)

    assert response.status_code == 200
    assert "attachment" in response.headers["content-disposition"]
    book = load_workbook(io.BytesIO(response.content))
    orders, items = book["Заказы"], book["Позиции"]
    # Время — московское: 9:30 UTC → 12:30
    first = [cell.value for cell in orders[2]][:5]
    assert first == [1, datetime(2026, 10, 20, 12, 30), "Выдан", "Сергей", "+79991112233"]
    assert orders["H2"].value == "распил по 3 м"
    # Итог без отменённого заказа — формулой
    assert orders[f"I{orders.max_row}"].value == '=SUMIF(C2:C3,"<>Отменён",I2:I3)'
    assert [cell.value for cell in items[2]][2:8] == ["Доска 50×150", "шт", 3, "шт", 620, 1860]


def test_export_is_for_admins_only(client: TestClient):
    setup_orders()
    customer = {"Authorization": f"Bearer {create_access_token(user_id=2, role='CUSTOMER')}"}

    assert client.get("/api/v1/admin/orders/export", headers=customer).status_code == 403
